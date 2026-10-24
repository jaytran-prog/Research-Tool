"""PowerPoint backend: draws layout primitives as native shapes, tables and charts (all editable)."""
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from deck_layout import W, H, cell_text, fmt_value

PX = 9525  # EMU per pixel at 96 dpi
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def E_(v):
    return Emu(int(round(v * PX)))


def pt(px):
    return Pt(px * 0.75)


def spacing(lh):
    """PowerPoint's single spacing is already ~1.2x the font size; convert a CSS-style line height."""
    return round(lh / 1.2, 2)


def drop_theme_style(shape):
    """Remove the theme style reference (default blue fill, outline and shadow); we set everything explicitly."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


class Painter:
    def __init__(self, theme):
        self.theme = theme

    def rgb(self, c):
        c = c if c.startswith("#") else self.theme["colors"].get(c, "#000000")
        return RGBColor.from_string(c.lstrip("#").upper())

    def font_name(self, key):
        return self.theme["fonts"]["head" if key == "head" else "body"]

    def _style_shape(self, shp, fill, line, line_w, hatch=False):
        drop_theme_style(shp)
        shp.shadow.inherit = False
        if hatch and fill:
            shp.fill.patterned()
            shp.fill.pattern = MSO_PATTERN_TYPE.WIDE_UPWARD_DIAGONAL
            shp.fill.fore_color.rgb = self.rgb(fill)
            shp.fill.back_color.rgb = self.rgb("white")
            line, line_w = fill, 1.5
        elif fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = self.rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = self.rgb(line)
            shp.line.width = pt(line_w)
        else:
            shp.line.fill.background()

    def text(self, slide, p):
        tb = slide.shapes.add_textbox(E_(p["x"]), E_(p["y"]), E_(max(p["w"], 1)), E_(max(p["h"], 1)))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = ANCHOR[p["valign"]]
        para = tf.paragraphs[0]
        para.alignment = ALIGN[p["align"]]
        para.line_spacing = spacing(p["lh"])
        for r in p["runs"]:
            for k, piece in enumerate(r.get("text", "").split("\n")):
                if k > 0:
                    para = tf.add_paragraph()
                    para.alignment = ALIGN[p["align"]]
                    para.line_spacing = spacing(p["lh"])
                if not piece:
                    continue
                run = para.add_run()
                run.text = piece
                f = run.font
                f.name = self.font_name(p["font"])
                f.size = pt(r.get("size") or p["size"])
                f.bold = r["bold"] if r.get("bold") is not None else p["bold"]
                f.italic = bool(r.get("italic") or p["italic"])
                f.color.rgb = self.rgb(r.get("color") or p["color"])

    def rect(self, slide, p, shape=MSO_SHAPE.RECTANGLE):
        radius = p.get("radius", 0)
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if (radius and shape == MSO_SHAPE.RECTANGLE) else shape
        shp = slide.shapes.add_shape(kind, E_(p["x"]), E_(p["y"]), E_(max(p["w"], 1)), E_(max(p["h"], 1)))
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = min(0.5, radius / max(1, min(p["w"], p["h"])))
        self._style_shape(shp, p.get("fill"), p.get("line"), p.get("line_w", 1), p.get("hatch"))
        return shp

    def line(self, slide, p):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E_(p["x1"]), E_(p["y1"]), E_(p["x2"]), E_(p["y2"]))
        drop_theme_style(c)
        c.line.color.rgb = self.rgb(p["color"])
        c.line.width = pt(p["w"])
        if p["dash"]:
            c.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    def diamond(self, slide, p):
        s = 0.8
        q = dict(p, x=p["x"] + p["w"] * (1 - s) / 2, y=p["y"] + p["h"] * (1 - s) / 2, w=p["w"] * s, h=p["h"] * s)
        self.rect(slide, q, MSO_SHAPE.DIAMOND)

    def harvey(self, slide, p):
        x, y, d, score = p["x"], p["y"], p["d"], p["score"]
        if score >= 4:
            self.rect(slide, {"x": x, "y": y, "w": d, "h": d, "fill": "ink"}, MSO_SHAPE.OVAL)
            return
        if score > 0:
            pie = slide.shapes.add_shape(MSO_SHAPE.PIE, E_(x), E_(y), E_(d), E_(d))
            pie.adjustments[0] = 270 * 0.6
            pie.adjustments[1] = ((270 + 90 * score) % 360) * 0.6
            self._style_shape(pie, "ink", None, 0)
        self.rect(slide, {"x": x, "y": y, "w": d, "h": d, "fill": None, "line": "ink", "line_w": 2}, MSO_SHAPE.OVAL)

    def table(self, slide, p):
        rows, cols = len(p["rows"]), len(p["col_w"])
        gf = slide.shapes.add_table(rows, cols, E_(p["x"]), E_(p["y"]), E_(p["w"]), E_(p["h"]))
        tbl_pr = gf._element.graphic.graphicData.tbl.tblPr
        sid = tbl_pr.find(qn("a:tableStyleId"))
        if sid is None:
            sid = etree.SubElement(tbl_pr, qn("a:tableStyleId"))
        sid.text = NO_STYLE_TABLE
        tbl = gf.table
        tbl.first_row = False
        tbl.horz_banding = False
        for c, w in enumerate(p["col_w"]):
            tbl.columns[c].width = E_(w)
        for r, h in enumerate(p["row_h"]):
            tbl.rows[r].height = E_(h)
        for r, row in enumerate(p["rows"]):
            fill = p["fills"][r]
            for c in range(cols):
                cell = tbl.cell(r, c)
                raw = row[c] if c < len(row) else ""
                if fill:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = self.rgb(fill)
                else:
                    cell.fill.background()
                cell.margin_left = cell.margin_right = E_(10)
                cell.margin_top = cell.margin_bottom = E_(6)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                para = tf.paragraphs[0]
                para.line_spacing = spacing(1.3)
                run = para.add_run()
                run.text = cell_text(raw)
                f = run.font
                f.name = self.font_name("body")
                f.size = pt(p["size"])
                f.bold = bool(isinstance(raw, dict) and raw.get("bold"))
                f.color.rgb = self.rgb(raw.get("color") if isinstance(raw, dict) and raw.get("color") else "ink")

    def bar(self, slide, p):
        spec = p["spec"]
        cats, series = spec.get("categories", []), spec.get("series", [])
        if not cats or not series:
            return
        horizontal = spec.get("orientation", "horizontal") == "horizontal"
        highlight = set(spec.get("highlight", []))
        # Match the HTML backend: all-negative data is charted as magnitudes with signed labels.
        mirror = max(v for ser in series for v in ser.get("values", []) if v is not None) <= 0
        data = CategoryChartData()
        data.categories = cats
        for ser in series:
            vals = [abs(v) if (mirror and v is not None) else v for v in ser.get("values", [])]
            data.add_series(ser.get("name", "") + (" (magnitude)" if mirror else ""), vals)
        kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
        chart = slide.shapes.add_chart(kind, E_(p["x"]), E_(p["y"]), E_(p["w"]), E_(p["h"]), data).chart
        chart.font.name = self.font_name("body")
        chart.font.size = pt(14)
        chart.font.color.rgb = self.rgb("ink")
        chart.has_title = False
        chart.has_legend = len(series) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False
        plot = chart.plots[0]
        plot.gap_width = 55
        plot.overlap = -10 if len(series) > 1 else 0
        va = chart.value_axis
        va.visible = False
        va.has_major_gridlines = False
        va.has_minor_gridlines = False
        ca = chart.category_axis
        ca.has_major_gridlines = False
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        ca.format.line.color.rgb = self.rgb("grey_light")
        if horizontal:
            ca.reverse_order = True
        palette = ["accent", "ink", "grey", "grey_light"]
        n = len(cats)
        for j, s in enumerate(chart.series):
            src = series[j]
            ests = src.get("estimate") or [False] * n
            if len(series) > 1:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = self.rgb(palette[j % 4])
            for i in range(n):
                v = src["values"][i]
                if v is None:
                    continue
                if len(series) == 1:
                    color = "accent" if (i in highlight or not highlight) else "grey_bar"
                else:
                    color = palette[j % 4]
                point = s.points[i]
                if ests[i]:
                    point.format.fill.patterned()
                    point.format.fill.pattern = MSO_PATTERN_TYPE.WIDE_UPWARD_DIAGONAL
                    point.format.fill.fore_color.rgb = self.rgb(color)
                    point.format.fill.back_color.rgb = self.rgb("white")
                    point.format.line.color.rgb = self.rgb(color)
                else:
                    point.format.fill.solid()
                    point.format.fill.fore_color.rgb = self.rgb(color)
                dl = point.data_label
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
                tf = dl.text_frame
                tf.text = fmt_value(v, spec) + (" est." if ests[i] else "")
                for para in tf.paragraphs:
                    for run in para.runs:
                        run.font.size = pt(14)
                        run.font.bold = i in highlight
                        run.font.color.rgb = self.rgb("ink")


def render_pptx(built, deck, theme, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = E_(W), E_(H)
    blank = prs.slide_layouts[6]
    painter = Painter(theme)
    for s in built:
        slide = prs.slides.add_slide(blank)
        for p in s["prims"]:
            t = p["type"]
            if t == "text":
                painter.text(slide, p)
            elif t == "rect":
                painter.rect(slide, p)
            elif t == "ellipse":
                painter.rect(slide, p, MSO_SHAPE.OVAL)
            elif t == "line":
                painter.line(slide, p)
            elif t == "diamond":
                painter.diamond(slide, p)
            elif t == "harvey":
                painter.harvey(slide, p)
            elif t == "table":
                painter.table(slide, p)
            elif t == "bar":
                painter.bar(slide, p)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    meta = deck.get("meta", {})
    prs.core_properties.title = meta.get("title", "")
    prs.core_properties.author = meta.get("prepared_by", "")
    prs.save(out_path)
