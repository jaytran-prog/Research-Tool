#!/usr/bin/env python3
"""Validate a written report spec and render it as a web article (HTML), PDF and/or Word document.

Usage:
  python3 render_report.py report.json [--formats html,pdf,docx] [--theme default|path.json] [--out DIR] [--check]
"""
import argparse
import html
import json
import os
import re
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from deck_layout import EVIDENCE_LABEL, EXHIBITS, exhibit  # noqa: E402
from render_deck import LEVELS, NEEDS_SOURCES, check_exhibit, find_chrome, html_to_pdf, load_theme  # noqa: E402
from render_html import _c, _prim  # noqa: E402

FIG_W = 650
BLOCKS = {"p", "exhibit", "callout", "bullets"}
CITE = re.compile(r"\s*\[(\d+(?:\s*,\s*\d+)*)\]")


# ---------- validation ----------

def prose_words(rep):
    texts = list(rep.get("at_a_glance", [])) + [rep.get("about", "")]
    for s in rep.get("sections", []):
        texts.append(s.get("heading", ""))
        for b in s.get("blocks", []):
            if b.get("type") in ("p", "callout"):
                texts.append(b.get("text", ""))
            elif b.get("type") == "bullets":
                texts += b.get("items", [])
    return sum(len(CITE.sub("", t).split()) for t in texts)


def validate(rep):
    errors, warnings = [], []
    meta = rep.get("meta", {})
    if not meta.get("title"):
        errors.append("meta.title is missing")
    source_ns = {s.get("n") for s in rep.get("sources", [])}
    sections = rep.get("sections", [])
    if len(sections) < 3:
        warnings.append("fewer than 3 sections")
    if not 3 <= len(rep.get("at_a_glance", [])) <= 5:
        warnings.append("at_a_glance should have 3–5 bullets")

    def cites(text, where):
        for group in CITE.findall(text or ""):
            for n in (int(x) for x in group.split(",")):
                if n not in source_ns:
                    errors.append(f"{where}: citation [{n}] is not in sources")

    for g in rep.get("at_a_glance", []):
        cites(g, "at_a_glance")
    cites(rep.get("about", ""), "about")
    n_ex = 0
    for i, s in enumerate(sections, 1):
        where = f"section {i}"
        if not s.get("heading"):
            errors.append(f"{where}: missing heading")
        for b in s.get("blocks", []):
            t = b.get("type")
            if t not in BLOCKS:
                errors.append(f"{where}: unknown block type '{t}'")
                continue
            if t in ("p", "callout"):
                cites(b.get("text", ""), where)
            elif t == "bullets":
                for item in b.get("items", []):
                    cites(item, where)
            elif t == "exhibit":
                n_ex += 1
                ex = b.get("exhibit") or {}
                kind = ex.get("kind")
                w = f"{where}, exhibit {n_ex}"
                if kind not in EXHIBITS or kind == "hero_stat":
                    errors.append(f"{w}: exhibit kind '{kind}' is not allowed in reports")
                    continue
                if not b.get("title"):
                    errors.append(f"{w}: missing title")
                if b.get("evidence") not in LEVELS:
                    errors.append(f"{w}: evidence must be strong, mixed or weak")
                if kind in NEEDS_SOURCES and not b.get("sources"):
                    errors.append(f"{w}: {kind} shows data but cites no sources")
                for n in b.get("sources", []) or []:
                    if n not in source_ns:
                        errors.append(f"{w}: source [{n}] is not in sources")
                errors += [f"{w}: {m}" for m in check_exhibit(kind, ex)]
    if not 4 <= n_ex <= 7:
        warnings.append(f"{n_ex} exhibits (aim for 4–7)")
    words = prose_words(rep)
    if not 1500 <= words <= 3500:
        warnings.append(f"about {words} words of prose (aim for 1,800–3,000)")
    return errors, warnings, words


# ---------- shared pieces ----------

def numbered_exhibits(rep):
    """Yields (section_index, block_index, exhibit_number, block) for every exhibit in order."""
    n = 0
    for si, s in enumerate(rep.get("sections", [])):
        for bi, b in enumerate(s.get("blocks", [])):
            if b.get("type") == "exhibit":
                n += 1
                yield si, bi, n, b


def fig_height(b):
    return max(200, min(520, int(b.get("height", 320))))


def fig_html(b, theme):
    h = fig_height(b)
    prims = exhibit((0, 0, FIG_W, h), b.get("exhibit"))
    return f'<div class="fig" style="height:{h}px">' + "".join(_prim(p, theme) for p in prims) + "</div>"


def source_line(b, by_n):
    parts = [f"[{n}] {by_n[n].get('label', '')}" for n in b.get("sources", []) or [] if n in by_n]
    return ("Source: " + "; ".join(parts)) if parts else ""


# ---------- HTML / PDF ----------

CSS = """
*{box-sizing:border-box;margin:0;padding:0}
body{background:#fff;color:%(ink)s;font-family:%(body)s}
.rep{width:%(w)spx;margin:48px auto 80px}
.kicker{font-size:12px;font-weight:700;letter-spacing:.06em;color:%(accent)s;text-transform:uppercase}
.bar{width:72px;height:5px;background:%(accent)s;margin:18px 0 20px}
h1{font-family:%(head)s;font-weight:400;font-size:38px;line-height:1.15;margin-bottom:16px}
.dek{font-size:19px;line-height:1.45;color:%(grey_dark)s;margin-bottom:18px}
.meta{font-size:13px;color:%(grey_dark)s}
.note{font-size:12px;color:%(grey)s;font-style:italic;margin-top:6px;line-height:1.45}
.badge{display:inline-flex;align-items:center;gap:6px;font-size:12px;font-weight:700;color:%(grey)s}
.dot{width:9px;height:9px;border-radius:50%%;display:inline-block}
.glance{background:%(accent_soft)s;border-left:5px solid %(accent)s;padding:20px 24px;margin:30px 0 8px}
.glance h3,.callout .lab{font-size:12px;letter-spacing:.06em;text-transform:uppercase;color:%(accent)s;margin-bottom:10px}
.glance li{margin:0 0 8px 18px;font-size:16px;line-height:1.45}
h2{font-family:%(head)s;font-weight:400;font-size:25px;line-height:1.25;margin:40px 0 14px}
p,.bul li{font-size:16.5px;line-height:1.65;margin-bottom:14px}
.bul{margin:0 0 14px 22px}
sup a{color:%(accent)s;text-decoration:none;font-size:11px;font-weight:700}
.callout{border-left:4px solid %(accent)s;background:%(bg_soft)s;padding:16px 20px;margin:18px 0;break-inside:avoid}
.glance{break-inside:avoid}
.callout p{margin:0;font-size:16px}
figure.ex{margin:26px 0 24px;padding-top:12px;border-top:2px solid %(ink)s;break-inside:avoid}
.ex-k{font-size:11px;font-weight:700;letter-spacing:.08em;color:%(accent)s;margin-bottom:4px}
.ex-t{font-size:17px;font-weight:700;line-height:1.35;margin-bottom:14px}
.fig{position:relative;width:%(w)spx}
.t{position:absolute;display:flex;flex-direction:column;white-space:normal;overflow-wrap:anywhere}
.s{position:absolute}.l{position:absolute;left:0;top:0;pointer-events:none;overflow:visible}
.tb{position:absolute;border-collapse:collapse;table-layout:fixed}.tb td{padding:8px 10px;vertical-align:middle;line-height:1.3}
figcaption{font-size:11.5px;color:%(grey)s;line-height:1.45;margin-top:10px;display:flex;justify-content:space-between;gap:16px}
.src ol{margin-left:22px}.src li{font-size:12.5px;line-height:1.5;color:%(grey_dark)s;margin-bottom:6px;overflow-wrap:anywhere}
@page{size:A4;margin:18mm 0 16mm 0;@bottom-center{content:counter(page);font-size:9px;color:%(grey)s}}
@media print{.rep{margin:0 auto}h2{break-after:avoid}}
"""


def cite_html(text):
    def sub(m):
        nums = [x.strip() for x in m.group(1).split(",")]
        return "<sup>" + ",".join(f'<a href="#src-{n}">{n}</a>' for n in nums) + "</sup>"
    return CITE.sub(sub, html.escape(text))


def badge_html(level, theme):
    if level not in EVIDENCE_LABEL:
        return ""
    return f'<span class="badge"><span class="dot" style="background:{_c(theme, level)}"></span>{EVIDENCE_LABEL[level]}</span>'


def render_report_html(rep, theme):
    m = rep.get("meta", {})
    col = theme["colors"]
    vals = {"w": FIG_W, "ink": col["ink"], "accent": col["accent"], "accent_soft": col["accent_soft"],
            "grey": col["grey"], "grey_dark": col["grey_dark"], "bg_soft": col["bg_soft"],
            "body": theme["fonts"]["body_css"], "head": theme["fonts"]["head_css"]}
    by_n = {s.get("n"): s for s in rep.get("sources", [])}
    out = [f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" '
           f'content="width=device-width,initial-scale=1"><title>{html.escape(m.get("title", "Report"))}</title>'
           f"<style>{CSS % vals}</style></head><body><article class=\"rep\"><header>"]
    if m.get("prepared_by"):
        out.append(f'<div class="kicker">{html.escape(m["prepared_by"])}</div>')
    out.append(f'<div class="bar"></div><h1>{html.escape(m.get("title", ""))}</h1>')
    if m.get("dek"):
        out.append(f'<div class="dek">{html.escape(m["dek"])}</div>')
    line = " · ".join(v for v in (m.get("client"), m.get("date")) if v)
    out.append(f'<div class="meta">{html.escape(line)} &nbsp; {badge_html(m.get("evidence"), theme)}</div>')
    if m.get("note"):
        out.append(f'<div class="note">{html.escape(m["note"])}</div>')
    out.append("</header>")
    if rep.get("at_a_glance"):
        out.append('<aside class="glance"><h3>At a glance</h3><ul>')
        out += [f"<li>{cite_html(g)}</li>" for g in rep["at_a_glance"]]
        out.append("</ul></aside>")
    ex_num = {(si, bi): n for si, bi, n, _ in numbered_exhibits(rep)}
    for si, s in enumerate(rep.get("sections", [])):
        out.append(f'<section><h2>{html.escape(s.get("heading", ""))}</h2>')
        for bi, b in enumerate(s.get("blocks", [])):
            t = b.get("type")
            if t == "p":
                out.append(f"<p>{cite_html(b.get('text', ''))}</p>")
            elif t == "bullets":
                out.append('<ul class="bul">' + "".join(f"<li>{cite_html(i)}</li>" for i in b.get("items", [])) + "</ul>")
            elif t == "callout":
                lab = html.escape(b.get("label", "What this means"))
                out.append(f'<div class="callout"><div class="lab">{lab}</div><p>{cite_html(b.get("text", ""))}</p></div>')
            elif t == "exhibit":
                cap = " · ".join(x for x in (b.get("note", ""), source_line(b, by_n)) if x)
                out.append(f'<figure class="ex"><div class="ex-k">EXHIBIT {ex_num[(si, bi)]}</div>'
                           f'<div class="ex-t">{html.escape(b.get("title", ""))}</div>{fig_html(b, theme)}'
                           f'<figcaption><span>{html.escape(cap)}</span>{badge_html(b.get("evidence"), theme)}'
                           f"</figcaption></figure>")
        out.append("</section>")
    if rep.get("about"):
        out.append("<section><h2>About this research</h2>")
        out += [f"<p>{cite_html(para)}</p>" for para in rep["about"].split("\n\n")]
        out.append("</section>")
    if rep.get("sources"):
        out.append('<section class="src"><h2>Sources</h2><ol>')
        for s in rep["sources"]:
            url = html.escape(s.get("url", ""))
            out.append(f'<li id="src-{s.get("n")}" value="{s.get("n")}">{html.escape(s.get("label", ""))} — '
                       f'<a href="{url}">{url}</a></li>')
        out.append("</ol></section>")
    out.append("</article></body></html>")
    return "".join(out)


# ---------- chart images for Word ----------

def exhibit_pngs(rep, theme, workdir):
    """Renders every exhibit to a PNG (2x) with one headless Chrome screenshot, then crops. Returns {n: path}."""
    chrome = find_chrome()
    items = list(numbered_exhibits(rep))
    if not chrome or not items:
        return {}
    try:
        from PIL import Image
    except ImportError:
        return {}
    gap = 20
    offsets, y, body = [], 0, []
    for _, _, n, b in items:
        h = fig_height(b)
        body.append(f'<div style="position:absolute;left:0;top:{y}px">{fig_html(b, theme)}</div>')
        offsets.append((n, y, h))
        y += h + gap
    page = (f'<!doctype html><html><head><meta charset="utf-8"><style>*{{margin:0;padding:0;box-sizing:border-box}}'
            f'body{{background:#fff;font-family:{theme["fonts"]["body_css"]};color:{theme["colors"]["ink"]};'
            f'width:{FIG_W}px;height:{y}px;position:relative;overflow:hidden}}'
            '.fig{position:relative}.t{position:absolute;display:flex;flex-direction:column;overflow-wrap:anywhere}'
            '.s{position:absolute}.l{position:absolute;left:0;top:0;overflow:visible}'
            '.tb{position:absolute;border-collapse:collapse;table-layout:fixed}.tb td{padding:8px 10px;'
            f'vertical-align:middle;line-height:1.3}}</style></head><body>{"".join(body)}</body></html>')
    src = os.path.join(workdir, "figs.html")
    shot = os.path.join(workdir, "figs.png")
    with open(src, "w", encoding="utf-8") as f:
        f.write(page)
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run([chrome, "--headless=new", "--disable-gpu", "--hide-scrollbars", f"--user-data-dir={profile}",
                        "--force-device-scale-factor=2", f"--window-size={FIG_W},{y}", f"--screenshot={shot}",
                        "--virtual-time-budget=3000", "file://" + src],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
    if not os.path.exists(shot):
        return {}
    img = Image.open(shot)
    scale = img.width / FIG_W
    paths = {}
    for n, top, h in offsets:
        p = os.path.join(workdir, f"exhibit-{n}.png")
        img.crop((0, int(top * scale), img.width, int((top + h) * scale))).save(p)
        paths[n] = p
    return paths


# ---------- Word ----------

def render_report_docx(rep, theme, out_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    col = theme["colors"]
    head_font, body_font = theme["fonts"]["head"], theme["fonts"]["body"]

    def rgb(key):
        return RGBColor.from_string(_c(theme, key).lstrip("#").upper())

    def set_font(run, name=body_font, size=10.5, bold=False, italic=False, color="ink", sup=False):
        run.font.name = name
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = rgb(color)
        run.font.superscript = sup

    def para(text="", size=10.5, color="ink", bold=False, italic=False, font=body_font, after=6, before=0,
             cited=False, align=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.line_spacing = 1.25
        if align:
            p.alignment = align
        add_text(p, text, size, color, bold, italic, font, cited)
        return p

    def add_text(p, text, size=10.5, color="ink", bold=False, italic=False, font=body_font, cited=True):
        pos = 0
        for m in (CITE.finditer(text) if cited else []):
            if m.start() > pos:
                set_font(p.add_run(text[pos:m.start()]), font, size, bold, italic, color)
            set_font(p.add_run(m.group(1).replace(" ", "")), font, size, True, False, "accent", sup=True)
            pos = m.end()
        if pos < len(text):
            set_font(p.add_run(text[pos:]), font, size, bold, italic, color)

    def shaded_box(fill_key, bar_key):
        tbl = doc.add_table(rows=1, cols=1)
        cell = tbl.cell(0, 0)
        tcpr = cell._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), _c(theme, fill_key).lstrip("#"))
        tcpr.append(shd)
        borders = OxmlElement("w:tcBorders")
        for side in ("top", "bottom", "right"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "nil")
            borders.append(el)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "36")
        left.set(qn("w:color"), _c(theme, bar_key).lstrip("#"))
        borders.append(left)
        tcpr.append(borders)
        mar = OxmlElement("w:tcMar")
        for side in ("top", "bottom", "left", "right"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:w"), "200")
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tcpr.append(mar)
        cell.paragraphs[0].text = ""
        return cell

    def cell_para(cell, text, first=False, **kw):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(kw.pop("after", 4))
        p.paragraph_format.line_spacing = 1.25
        add_text(p, text, **kw)
        return p

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin, sec.bottom_margin = Mm(20), Mm(18)
    normal = doc.styles["Normal"]
    normal.font.name = body_font
    normal.font.size = Pt(10.5)

    m = rep.get("meta", {})
    by_n = {s.get("n"): s for s in rep.get("sources", [])}
    if m.get("prepared_by"):
        para(m["prepared_by"].upper(), size=9, color="accent", bold=True, after=10, cited=False)
    para(m.get("title", ""), size=26, font=head_font, after=10, cited=False)
    if m.get("dek"):
        para(m["dek"], size=13, color="grey_dark", after=10, cited=False)
    line = " · ".join(v for v in (m.get("client"), m.get("date")) if v)
    if m.get("evidence") in EVIDENCE_LABEL:
        line += f"   ●  {EVIDENCE_LABEL[m['evidence']]}"
    para(line, size=9, color="grey_dark", after=4, cited=False)
    if m.get("note"):
        para(m["note"], size=8.5, color="grey", italic=True, after=12, cited=False)

    if rep.get("at_a_glance"):
        cell = shaded_box("accent_soft", "accent")
        cell_para(cell, "AT A GLANCE", first=True, size=9, color="accent", bold=True, cited=False, after=6)
        for g in rep["at_a_glance"]:
            cell_para(cell, "•  " + g, size=10.5)
        para("", after=4)

    pngs = {}
    with tempfile.TemporaryDirectory() as work:
        pngs = exhibit_pngs(rep, theme, work)
        ex_num = {(si, bi): n for si, bi, n, _ in numbered_exhibits(rep)}
        for si, s in enumerate(rep.get("sections", [])):
            h = para(s.get("heading", ""), size=16, font=head_font, before=14, after=8, cited=False)
            h.paragraph_format.keep_with_next = True
            for bi, b in enumerate(s.get("blocks", [])):
                t = b.get("type")
                if t == "p":
                    para(b.get("text", ""), cited=True, after=8)
                elif t == "bullets":
                    for item in b.get("items", []):
                        p = doc.add_paragraph(style="List Bullet")
                        p.paragraph_format.space_after = Pt(4)
                        add_text(p, item)
                elif t == "callout":
                    cell = shaded_box("bg_soft", "accent")
                    cell_para(cell, b.get("label", "What this means").upper(), first=True, size=9, color="accent",
                              bold=True, cited=False, after=4)
                    cell_para(cell, b.get("text", ""), size=10.5)
                    para("", after=2)
                elif t == "exhibit":
                    n = ex_num[(si, bi)]
                    k = para(f"EXHIBIT {n}", size=8.5, color="accent", bold=True, before=8, after=2, cited=False)
                    k.paragraph_format.keep_with_next = True
                    tp = para(b.get("title", ""), size=11, bold=True, after=6, cited=False)
                    tp.paragraph_format.keep_with_next = True
                    if n in pngs:
                        pic = doc.add_paragraph()
                        pic.paragraph_format.keep_with_next = True
                        pic.add_run().add_picture(pngs[n], width=Mm(170))
                    else:
                        para("[Chart: see the PDF or web version of this report]", size=9, color="grey", italic=True,
                             cited=False)
                    cap = " · ".join(x for x in (b.get("note", ""), source_line(b, by_n)) if x)
                    if b.get("evidence") in EVIDENCE_LABEL:
                        cap += ("   ●  " if cap else "●  ") + EVIDENCE_LABEL[b["evidence"]]
                    para(cap, size=8, color="grey", after=12, cited=False)

    if rep.get("about"):
        para("About this research", size=16, font=head_font, before=14, after=8, cited=False)
        for chunk in rep["about"].split("\n\n"):
            para(chunk, cited=True, after=8)
    if rep.get("sources"):
        para("Sources", size=16, font=head_font, before=14, after=8, cited=False)
        for s in rep["sources"]:
            para(f"[{s.get('n')}] {s.get('label', '')} — {s.get('url', '')}", size=8.5, color="grey_dark", after=4,
                 cited=False)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run(m.get("title", "") + "   ·   "), size=8, color="grey")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    fld.append(run)
    footer._p.append(fld)
    doc.core_properties.title = m.get("title", "")
    doc.core_properties.author = m.get("prepared_by", "")
    doc.save(out_path)
    return bool(pngs)


# ---------- CLI ----------

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("report")
    ap.add_argument("--formats", default="html,pdf,docx")
    ap.add_argument("--theme", default="default")
    ap.add_argument("--out", default=None)
    ap.add_argument("--check", action="store_true", help="validate only")
    args = ap.parse_args()

    with open(args.report, encoding="utf-8") as f:
        rep = json.load(f)
    errors, warnings, words = validate(rep)
    for w in warnings:
        print("WARNING:", w)
    for e in errors:
        print("ERROR:", e)
    if errors:
        sys.exit(f"{len(errors)} error(s) — fix report.json and run again")
    if args.check:
        print(f"OK: report valid, about {words} words of prose")
        return

    theme = load_theme(args.theme)
    out_dir = args.out or os.path.dirname(os.path.abspath(args.report))
    os.makedirs(out_dir, exist_ok=True)
    stem = os.path.splitext(os.path.basename(args.report))[0]
    formats = [x.strip() for x in args.formats.split(",") if x.strip()]
    written = []
    html_path = os.path.join(out_dir, stem + ".html")
    if "html" in formats or "pdf" in formats:
        with open(html_path, "w", encoding="utf-8") as f:
            f.write(render_report_html(rep, theme))
        if "html" in formats:
            written.append(html_path)
    if "pdf" in formats:
        pdf_path = os.path.join(out_dir, stem + ".pdf")
        if html_to_pdf(html_path, pdf_path):
            written.append(pdf_path)
        else:
            print("WARNING: no Chrome/Chromium found — open the HTML and print to PDF instead")
        if "html" not in formats:
            os.remove(html_path)
    if "docx" in formats:
        try:
            import docx  # noqa: F401
        except ImportError:
            print("WARNING: python-docx is not installed — run: pip install python-docx")
        else:
            docx_path = os.path.join(out_dir, stem + ".docx")
            if not render_report_docx(rep, theme, docx_path):
                print("WARNING: charts could not be rendered as images (needs Chrome and Pillow); "
                      "the Word file has placeholders instead")
            written.append(docx_path)
    print(f"Rendered report (about {words} words):")
    for w in written:
        print("  " + w)


if __name__ == "__main__":
    main()
